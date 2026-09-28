import sys
import os

HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, "..", "_template"))
sys.path.insert(0, HERE)
import review_docx_builder as b
import content as c
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

OUT = os.path.join(HERE, "recensione-mylotteriesplay-2026.docx")

doc = b.new_document()
b.set_header_footer(doc, c.HEADER)
b.add_title(doc, c.TITLE, c.SUBTITLE)
b.add_rating_box(doc, *c.RATING)
b.add_cta_bar(doc, c.CTA)
b.add_divider(doc)

for block in c.BLOCKS:
    kind, args = block[0], block[1:]
    if kind == "h1":
        b.add_h1(doc, args[0])
    elif kind == "h2":
        b.add_h2(doc, args[0])
    elif kind == "body":
        b.add_body(doc, args[0])
    elif kind == "small":
        b.add_body(doc, args[0], size=9)
    elif kind == "info":
        b.add_info_table(doc, args[0])
    elif kind == "callout":
        b.add_callout_box(doc, args[0], args[1])
    elif kind == "table":
        b.add_data_table(doc, args[0], args[1], highlight_col=args[2])
    elif kind == "avg":
        b.add_average_score(doc, args[0])
    elif kind == "note":
        b.add_note_box(doc, args[0], args[1])
    elif kind == "final":
        b.add_final_score_box(doc, args[0])
    elif kind == "divider":
        b.add_divider(doc)

footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = b.Pt(14)
b._add_run(footer_p, c.FOOTER, italic=True, color=b.GRAY, size=8)

# Nessuna riga di tabella spezzata a cavallo di due pagine (fix di impaginazione introdotto su Goldbet)
for table in doc.tables:
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

doc.save(OUT)
print("Salvato:", OUT)
